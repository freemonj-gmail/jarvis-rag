"""Document extraction and ingestion for Jarvis.

Houses the per-format readers (zip, Excel, PowerPoint, image OCR, PDF text +
tables + OCR fallback, and Docling/Unstructured office parsing), the
``file_extractor`` extension -> reader map, and the directory loader with
per-file (path, mtime, size) extraction caching.
"""

import hashlib
import html as html_lib
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from typing import Any, Dict, List, Optional

import cv2
import easyocr
import numpy as np
import pandas as pd
import pytesseract
from PIL import Image
from pptx import Presentation

from llama_index.core.schema import Document
from llama_index.readers.file.unstructured import UnstructuredReader

from jarvisrag.config import (
    JARVIS_APPLE_OCR_LANGS,
    JARVIS_APPLE_OCR_LEVEL,
    JARVIS_DEBUG,
    JARVIS_ENABLE_APPLE_OCR,
    JARVIS_ENABLE_PAGE_PIPELINE,
    JARVIS_ENABLE_SURYA,
    JARVIS_ENABLE_VLM,
    JARVIS_OCR_ESCALATE_ON_FAILURE,
    JARVIS_OCR_MIN_CONF,
    JARVIS_PDF_MIN_CHARS_PER_PAGE,
    JARVIS_PDF_OCR_DPI,
    JARVIS_VLM_CAPTION,
    JARVIS_VLM_MODEL,
    logger,
)
from jarvisrag.iotools import suppress_java_warnings, time_limit
from jarvisrag.text.textproc import normalize_text
from jarvisrag.extraction.indexing import chunk_text
from jarvisrag.extraction.pagepipe import (
    Page,
    PageProcessor,
    open_pdf,
    pymupdf_available,
)


# Surya / TrOCR optional imports
try:
    from surya.ocr import run_ocr  # type: ignore[import-not-found]
    from surya.model.detection import load_model as load_detection_model  # type: ignore[import-not-found]
    from surya.model.recognition import load_model as load_recognition_model  # type: ignore[import-not-found]
    surya_available = True
except Exception:
    surya_available = False

try:
    from transformers import TrOCRProcessor, VisionEncoderDecoderModel
    trocr_available = True
except Exception:
    trocr_available = False


def _documents_from_page(page: "Page", file_path_str: str, file_name: str) -> List["Document"]:
    """Chunk a page-pipeline :class:`Page` into Documents in reading order.

    Page-level provenance (which engine produced the text) and the mean
    recognition confidence are recorded in metadata so later phases can prefer
    high-confidence native text over rescued OCR when ranking.
    """
    text = page.text()
    if not text:
        return []
    confidence = round(page.mean_confidence(), 4)
    docs: List[Document] = []
    for chunk_index, chunk in enumerate(chunk_text(text)):
        metadata = {
            "file_path": file_path_str,
            "file_name": file_name,
            "chunk_index": chunk_index,
            "extracted_by": page.source,
            "ocr_confidence": confidence,
        }
        if page.page_number is not None:
            metadata["page_number"] = page.page_number
        docs.append(Document(text=chunk, metadata=metadata))
    return docs


class ZipExtractor:
    def __init__(self, file_extractors):
        self.file_extractors = file_extractors

    def load_data(self, file_path, **kwargs):
        docs = []
        with zipfile.ZipFile(file_path, 'r') as zip_ref:
            with tempfile.TemporaryDirectory() as tmpdirname:
                zip_ref.extractall(tmpdirname)
                for root, _, files in os.walk(tmpdirname):
                    for name in files:
                        ext = os.path.splitext(name)[1].lower()
                        extractor = self.file_extractors.get(ext)
                        if extractor:
                            extracted_docs = extractor.load_data(os.path.join(root, name), **kwargs)
                            docs.extend(extracted_docs)
        return docs


class ExcelReader:
    def load_data(self, file_path, **kwargs):
        docs = []
        file_name = os.path.basename(file_path)
        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name in xls.sheet_names:
                df = xls.parse(sheet_name)
                text = normalize_text(df.to_string(index=False))
                if not text:
                    continue
                for chunk_index, chunk in enumerate(chunk_text(text)):
                    docs.append(Document(text=chunk, metadata={
                        "file_path": str(file_path),
                        "file_name": file_name,
                        "sheet": sheet_name,
                        "chunk_index": chunk_index,
                        "extracted_by": "excel"
                    }))
        except Exception as e:
            logger.error(f"Failed to read Excel file {file_path}: {e}")
        return docs


class PowerPointReader:
    def load_data(self, file_path, **kwargs):
        docs = []
        file_name = os.path.basename(file_path)
        try:
            prs = Presentation(file_path)
            for slide_index, slide in enumerate(prs.slides, start=1):
                slide_text = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text += shape.text + "\n"
                slide_text = normalize_text(slide_text)
                if not slide_text:
                    continue
                for chunk_index, chunk in enumerate(chunk_text(slide_text)):
                    docs.append(Document(text=chunk, metadata={
                        "file_path": str(file_path),
                        "file_name": file_name,
                        "slide_index": slide_index,
                        "chunk_index": chunk_index,
                        "extracted_by": "powerpoint"
                    }))
        except Exception as e:
            logger.error(f"Failed to read PowerPoint file {file_path}: {e}")
        return docs


class JpegExtractor:

    # -------------------------------------------------
    # CLASS-LEVEL SHARED MODELS
    # -------------------------------------------------

    _easyocr_reader = None

    _surya_det_model = None
    _surya_rec_model = None

    _trocr_processor = None
    _trocr_model = None

    _models_initialized = False

    # -------------------------------------------------
    # INITIALIZATION
    # -------------------------------------------------

    def __init__(self):

        if not JpegExtractor._models_initialized:

            logger.info("Initializing OCR models (one-time startup)...")

            # -----------------------------------------
            # EasyOCR
            # -----------------------------------------
            try:
                JpegExtractor._easyocr_reader = easyocr.Reader(
                    ['en'],
                    gpu=False
                )

                logger.info("EasyOCR initialized.")

            except Exception as e:
                logger.warning(f"EasyOCR initialization failed: {e}")

            # -----------------------------------------
            # Surya
            # -----------------------------------------
            if surya_available:

                try:
                    JpegExtractor._surya_det_model = load_detection_model()

                    JpegExtractor._surya_rec_model = (
                        load_recognition_model()
                    )

                    logger.info("Surya OCR initialized.")

                except Exception as e:
                    logger.warning(
                        f"Surya initialization failed: {e}"
                    )

            # -----------------------------------------
            # TrOCR
            # -----------------------------------------
            if trocr_available:

                try:

                    logger.info(
                        "Loading TrOCR processor..."
                    )

                    JpegExtractor._trocr_processor = (
                        TrOCRProcessor.from_pretrained(
                            'microsoft/trocr-base-printed'
                        )
                    )

                    logger.info(
                        "Loading TrOCR model..."
                    )

                    JpegExtractor._trocr_model = (
                        VisionEncoderDecoderModel.from_pretrained(
                            'microsoft/trocr-base-printed'
                        )
                    )

                    logger.info("TrOCR initialized.")

                except Exception as e:
                    logger.warning(
                        f"TrOCR initialization failed: {e}"
                    )

            JpegExtractor._models_initialized = True

    # -------------------------------------------------
    # IMAGE PREPROCESSING
    # -------------------------------------------------

    def _preprocess_image(self, image):

        cv_image = cv2.cvtColor(
            np.array(image),
            cv2.COLOR_RGB2BGR
        )

        gray = cv2.cvtColor(
            cv_image,
            cv2.COLOR_BGR2GRAY
        )

        h, w = gray.shape[:2]

        max_dim = max(h, w)

        if max_dim < 1600:

            scale = 1600 / max_dim

            gray = cv2.resize(
                gray,
                None,
                fx=scale,
                fy=scale,
                interpolation=cv2.INTER_CUBIC
            )

        clahe = cv2.createCLAHE(
            clipLimit=2.0,
            tileGridSize=(8, 8)
        )

        enhanced = clahe.apply(gray)

        denoised = cv2.fastNlMeansDenoising(
            enhanced,
            h=3
        )

        blurred = cv2.medianBlur(
            denoised,
            3
        )

        thresh = cv2.adaptiveThreshold(
            blurred,
            255,
            cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
            cv2.THRESH_BINARY,
            11,
            2,
        )

        return Image.fromarray(thresh)

    # -------------------------------------------------
    # IMAGE LOAD
    # -------------------------------------------------

    def _load_image(self, file_path_str):

        try:

            image = Image.open(file_path_str)

            return image.convert("RGB")

        except Exception as e:

            logger.error(
                f"Unable to open image {file_path_str}: {e}"
            )

            return None

    # -------------------------------------------------
    # SURYA OCR
    # -------------------------------------------------

    def _ocr_with_surya(self, image):

        if (
            not surya_available
            or JpegExtractor._surya_det_model is None
            or JpegExtractor._surya_rec_model is None
        ):
            return None

        try:

            predictions = run_ocr(
                [image],
                [JpegExtractor._surya_det_model],
                [JpegExtractor._surya_rec_model],
                langs=['en']
            )

            if predictions and predictions[0]:

                text = "\n".join(
                    [line.text for line in predictions[0]]
                )

                return text

        except Exception as e:

            logger.error(f"Surya OCR failed: {e}")

        return None

    # -------------------------------------------------
    # TrOCR
    # -------------------------------------------------

    def _ocr_with_trocr(self, image):

        if (
            not trocr_available
            or JpegExtractor._trocr_processor is None
            or JpegExtractor._trocr_model is None
        ):
            return None

        try:

            pixel_values = (
                JpegExtractor._trocr_processor(
                    image,
                    return_tensors="pt"
                ).pixel_values
            )

            generated_ids = (
                JpegExtractor._trocr_model.generate(
                    pixel_values,
                    max_new_tokens=256
                )
            )

            generated_text = (
                JpegExtractor._trocr_processor.batch_decode(
                    generated_ids,
                    skip_special_tokens=True
                )[0]
            )

            return generated_text

        except Exception as e:

            logger.error(f"TrOCR failed: {e}")

        return None

    # -------------------------------------------------
    # OCR QUALITY CHECK
    # -------------------------------------------------

    def _is_valid_ocr_text(self, text: str) -> bool:

        if not text:
            return False

        normalized = normalize_text(text)

        if len(normalized) < 20:

            keywords = [
                "dob",
                "date",
                "birth",
                "license",
                "name",
                "expiry",
                "issue",
                "address",
            ]

            lower = normalized.lower()

            if (
                any(keyword in lower for keyword in keywords)
                and len(re.findall(r"\w+", normalized)) >= 2
            ):
                return True

            return False

        if re.fullmatch(r"[\W_]+", normalized):
            return False

        tokens = re.findall(r"\w+", normalized)

        return len(tokens) >= 3

    # -------------------------------------------------
    # PAGE-PIPELINE FALLBACK OCR
    # -------------------------------------------------

    def _pipeline_fallback_ocr(self, crop):
        """Re-OCR a low-confidence region crop for the page pipeline.

        Tries pytesseract (fast, exposes per-word confidence) then easyocr, and
        returns ``(text, confidence)`` with confidence normalized to 0.0-1.0.
        """
        try:
            processed = self._preprocess_image(crop)
        except Exception:
            processed = crop

        try:
            data = pytesseract.image_to_data(
                processed, output_type=pytesseract.Output.DICT
            )
            words: List[str] = []
            confs: List[float] = []
            for word, conf in zip(data.get("text", []), data.get("conf", [])):
                if word and word.strip():
                    words.append(word)
                    try:
                        c = float(conf)
                    except (TypeError, ValueError):
                        c = -1.0
                    if c >= 0:
                        confs.append(c / 100.0)
            text = normalize_text(" ".join(words))
            if text:
                return text, (sum(confs) / len(confs) if confs else 0.0)
        except Exception as e:
            logger.info(f"pytesseract fallback failed: {e}")

        try:
            reader = JpegExtractor._easyocr_reader
            if reader is not None:
                result = reader.readtext(np.array(processed), detail=1)
                texts = [r[1] for r in result if len(r) >= 2 and r[1]]
                confs = [float(r[2]) for r in result if len(r) >= 3]
                text = normalize_text("\n".join(texts))
                if text:
                    return text, (sum(confs) / len(confs) if confs else 0.0)
        except Exception as e:
            logger.info(f"easyocr fallback failed: {e}")

        return "", 0.0

    # -------------------------------------------------
    # MAIN LOAD
    # -------------------------------------------------

    def load_data(self, file_path, **kwargs):

        file_path_str = str(file_path)

        file_name = os.path.basename(file_path)

        documents = []

        image = self._load_image(file_path_str)

        if image is None:
            return documents

        # -----------------------------------------
        # CONFIDENCE-ROUTED PAGE PIPELINE (opt-in)
        # -----------------------------------------
        # When enabled, run Surya layout + reading order + per-region
        # recognition confidence, routing only the low-confidence regions to the
        # pytesseract/easyocr fallback below. Falls back to the legacy cascade if
        # the pipeline yields nothing or errors.
        if JARVIS_ENABLE_PAGE_PIPELINE:
            try:
                processor = PageProcessor(fallback_ocr=self._pipeline_fallback_ocr)
                page = processor.process_image(image)
                pipeline_docs = _documents_from_page(page, file_path_str, file_name)
                if pipeline_docs:
                    return pipeline_docs
                logger.info(
                    f"Page pipeline produced no text for {file_name}; "
                    f"falling back to the legacy OCR cascade."
                )
            except Exception as e:
                logger.warning(
                    f"Page pipeline failed for {file_name} ({e}); "
                    f"falling back to the legacy OCR cascade."
                )

        processed_image = self._preprocess_image(image)

        # -----------------------------------------
        # SURYA
        # -----------------------------------------

        text_surya = self._ocr_with_surya(
            processed_image
        )

        if text_surya and text_surya.strip():

            text = normalize_text(text_surya)

            for chunk_index, chunk in enumerate(
                chunk_text(text)
            ):

                documents.append(Document(
                    text=chunk,
                    metadata={
                        "file_path": file_path_str,
                        "file_name": file_name,
                        "chunk_index": chunk_index,
                        "extracted_by": "surya"
                    }
                ))

            return documents

        # -----------------------------------------
        # TrOCR
        # -----------------------------------------

        text_trocr = self._ocr_with_trocr(
            processed_image
        )

        if text_trocr and text_trocr.strip():

            if self._is_valid_ocr_text(text_trocr):

                text = normalize_text(text_trocr)

                for chunk_index, chunk in enumerate(
                    chunk_text(text)
                ):

                    documents.append(Document(
                        text=chunk,
                        metadata={
                            "file_path": file_path_str,
                            "file_name": file_name,
                            "chunk_index": chunk_index,
                            "extracted_by": "trocr"
                        }
                    ))

                return documents

            logger.info(
                f"TrOCR result for {file_name} "
                f"appears too short/low-quality; "
                f"falling back to additional OCR engines. "
                f"Output: {repr(text_trocr)}"
            )

        # -----------------------------------------
        # PYTESSERACT
        # -----------------------------------------

        pyt_docs = []

        try:

            text_pt = normalize_text(
                pytesseract.image_to_string(
                    processed_image
                )
            )

            if text_pt:

                for chunk_index, chunk in enumerate(
                    chunk_text(text_pt)
                ):

                    pyt_docs.append(Document(
                        text=chunk,
                        metadata={
                            "file_path": file_path_str,
                            "file_name": file_name,
                            "chunk_index": chunk_index,
                            "extracted_by": "pytesseract"
                        }
                    ))

        except Exception as e:

            logger.error(
                f"pytesseract failed on {file_path}: {e}"
            )

        # -----------------------------------------
        # EASYOCR
        # -----------------------------------------

        eo_docs = []

        try:

            image_array = np.array(processed_image)

            result = (
                JpegExtractor._easyocr_reader.readtext(
                    image_array,
                    detail=0
                )
            )

            text_eo = normalize_text(
                "\n".join(result)
            )

            if text_eo:

                for chunk_index, chunk in enumerate(
                    chunk_text(text_eo)
                ):

                    eo_docs.append(Document(
                        text=chunk,
                        metadata={
                            "file_path": file_path_str,
                            "file_name": file_name,
                            "chunk_index": chunk_index,
                            "extracted_by": "easyocr"
                        }
                    ))

        except Exception as e:

            logger.error(
                f"easyocr failed on {file_path}: {e}"
            )

        documents.extend(pyt_docs)
        documents.extend(eo_docs)

        return documents


class TablePDFReader:
    def __init__(self):
        self._tabula_available = None
        self._tabula = None

    def _java_available(self) -> bool:
        if shutil.which("java") is None:
            return False
        try:
            result = subprocess.run(
                ["java", "-version"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=5,
                check=False,
            )
            return result.returncode == 0
        except Exception:
            return False

    def _check_tabula_available(self):
        if self._tabula_available is not None:
            return self._tabula_available

        if not self._java_available():
            logger.info("Java/JDK is not installed; disabling tabula-py table extraction.")
            self._tabula_available = False
            return False

        try:
            import tabula
            self._tabula = tabula
            self._tabula_available = True
            return True
        except ImportError:
            logger.info("tabula-py is not installed; disabling tabula-py table extraction.")
            self._tabula_available = False
            return False
        except Exception as e:
            logger.info(f"tabula-py initialization failed: {e}. Disabling tabula-py table extraction.")
            self._tabula_available = False
            return False

    def _disable_tabula(self, reason: str):
        if self._tabula_available is not False:
            logger.info(f"Disabling tabula-py table extraction: {reason}")
        self._tabula_available = False
        self._tabula = None

    def load_data(self, file_path: str, **kwargs):
        docs = []
        file_name = os.path.basename(file_path)

        if not self._check_tabula_available():
            tabula = None
        else:
            tabula = self._tabula

        if tabula is not None:
            try:
                with suppress_java_warnings():
                    tables = tabula.read_pdf(file_path, pages='all', multiple_tables=True)
                for i, table in enumerate(tables):
                    if not table.empty:
                        table_text = table.to_string(index=False)
                        if table_text.strip():
                            docs.append(Document(text=table_text, metadata={
                                "file_path": file_path,
                                "file_name": file_name,
                                "page_number": None,
                                "table_index": i,
                                "extracted_by": "tabula"
                            }))
                if docs:
                    return docs
            except Exception as e:
                err = str(e).lower()
                if any(keyword in err for keyword in ["jvm dll not found", "jvm", "java", "jpype"]):
                    self._disable_tabula(
                        "Java/JVM is unavailable or incorrectly configured for tabula-py. "
                        "Install a valid JDK only if you need tabula table extraction."
                    )
                else:
                    logger.info(f"Note: tabula table extraction failed for {file_path}: {e}")

        try:
            with suppress_java_warnings():
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    for page_num, page in enumerate(pdf.pages, start=1):
                        page_tables = page.extract_tables()
                        for i, table in enumerate(page_tables):
                            if table:
                                table_text = pd.DataFrame(table).to_string(index=False)
                                if table_text.strip():
                                    docs.append(Document(text=table_text, metadata={
                                        "file_path": file_path,
                                        "file_name": file_name,
                                        "page_number": page_num,
                                        "table_index": i,
                                        "extracted_by": "pdfplumber_table"
                                    }))
        except Exception as e:
            logger.info(f"Note: pdfplumber table extraction failed for {file_path}: {e}")
        return docs


class CombinedPDFReader:
    def __init__(self):
        self.table_reader = TablePDFReader()
        self.easyocr_reader = None
        self.pdf_file_timeout_seconds = float(os.environ.get("JARVIS_PDF_FILE_TIMEOUT_SECONDS", "90"))

    def _get_easyocr_reader(self):
        if self.easyocr_reader is None:
            logger.info("Initializing PDF EasyOCR fallback model...")
            self.easyocr_reader = easyocr.Reader(['en'])
            logger.info("PDF EasyOCR fallback model initialized.")
        return self.easyocr_reader

    def _build_page_docs(self, text: str, metadata: dict, source: str):
        docs = []
        chunks = chunk_text(text)
        for chunk_index, chunk in enumerate(chunks):
            chunk_meta = metadata.copy()
            chunk_meta.update({
                "chunk_index": chunk_index,
                "source": source,
            })
            docs.append(Document(text=chunk, metadata=chunk_meta))
        return docs

    def _pipeline_fallback_ocr(self, crop):
        """easyocr-based fallback for the page pipeline; returns (text, conf)."""
        try:
            reader = self._get_easyocr_reader()
            result = reader.readtext(np.array(crop), detail=1)
            texts = [r[1] for r in result if len(r) >= 2 and r[1]]
            confs = [float(r[2]) for r in result if len(r) >= 3]
            text = normalize_text("\n".join(texts))
            if text:
                return text, (sum(confs) / len(confs) if confs else 0.0)
        except Exception as e:
            logger.info(f"PDF pipeline easyocr fallback failed: {e}")
        return "", 0.0

    def _load_with_pipeline(self, file_path: str, file_name: str, **kwargs):
        """Confidence-routed path: native PyMuPDF text per page, OCR only for
        image-only/sparse pages. Returns ``None`` if the PDF can't be opened
        (signalling the caller to use the legacy pdfplumber path)."""
        doc = open_pdf(file_path)
        if doc is None:
            return None
        all_docs: List[Document] = []
        processor = PageProcessor(fallback_ocr=self._pipeline_fallback_ocr)
        try:
            with time_limit(self.pdf_file_timeout_seconds, f"PDF pipeline for {file_name}"):
                for page_index in range(doc.page_count):
                    fitz_page = doc.load_page(page_index)
                    page = processor.process_pdf_page(fitz_page, page_number=page_index + 1)
                    all_docs.extend(_documents_from_page(page, file_path, file_name))
        except TimeoutError as e:
            logger.warning(f"{e}. Returning partial content for {file_name}.")
        except Exception as e:
            logger.info(f"PDF pipeline extraction error for {file_name}: {e}")
        finally:
            try:
                doc.close()
            except Exception:
                pass
        # Tables stay on pdfplumber/tabula — they extract tabular structure far
        # better than page OCR does.
        try:
            table_docs = self.table_reader.load_data(file_path, **kwargs)
            if table_docs:
                all_docs.extend(table_docs)
        except Exception as e:
            logger.info(f"Table extraction failed for {file_name}: {e}")
        return all_docs

    def load_data(self, file_path: str, **kwargs):
        import pdfplumber
        all_docs = []
        file_name = os.path.basename(file_path)
        started = time.time()
        logger.info(f"Reading PDF: {file_name}")

        # -----------------------------------------
        # CONFIDENCE-ROUTED PAGE PIPELINE (opt-in)
        # -----------------------------------------
        if JARVIS_ENABLE_PAGE_PIPELINE and pymupdf_available:
            try:
                pipeline_docs = self._load_with_pipeline(file_path, file_name, **kwargs)
                if pipeline_docs:
                    logger.info(
                        f"Finished PDF (pipeline): {file_name} -> "
                        f"{len(pipeline_docs)} chunks in {time.time() - started:.1f}s"
                    )
                    return pipeline_docs
                if pipeline_docs is not None:
                    logger.info(
                        f"Page pipeline produced no content for {file_name}; "
                        f"falling back to the legacy PDF path."
                    )
            except Exception as e:
                logger.warning(
                    f"Page pipeline failed for {file_name} ({e}); "
                    f"falling back to the legacy PDF path."
                )

        try:
            with time_limit(self.pdf_file_timeout_seconds, f"PDF extraction for {file_name}"):
                with suppress_java_warnings():
                    with pdfplumber.open(file_path) as pdf:
                        for page_num, page in enumerate(pdf.pages, start=1):
                            page_text = normalize_text(page.extract_text())
                            page_metadata = {
                                "file_path": file_path,
                                "file_name": file_name,
                                "page_number": page_num,
                            }
                            if page_text:
                                all_docs.extend(self._build_page_docs(page_text, {**page_metadata, "extracted_by": "pdfplumber"}, "pdfplumber"))
                    table_docs = self.table_reader.load_data(file_path, **kwargs)
                    if table_docs:
                        all_docs.extend(table_docs)
        except TimeoutError as e:
            logger.warning(f"{e}. Returning partial content for {file_name}.")
        except Exception as e:
            logger.info(f"pdfplumber extraction error: {e}")

        if not all_docs:
            logger.info(f"No text extracted from {file_name} via text/table extraction. Attempting OCR fallback...")
            try:
                with time_limit(self.pdf_file_timeout_seconds, f"OCR extraction for {file_name}"):
                    with suppress_java_warnings():
                        with pdfplumber.open(file_path) as pdf:
                            for page_num, page in enumerate(pdf.pages, start=1):
                                page_metadata = {
                                    "file_path": file_path,
                                    "file_name": file_name,
                                    "page_number": page_num,
                                }
                                try:
                                    page_image = page.to_image(resolution=300).original.convert("RGB")
                                    image_array = np.array(page_image)
                                    result = self._get_easyocr_reader().readtext(image_array, detail=0)
                                    page_text = "\n".join(result)
                                    if page_text.strip():
                                        all_docs.extend(self._build_page_docs(page_text, {**page_metadata, "extracted_by": "easyocr"}, "easyocr"))
                                except Exception as e:
                                    logger.info(f"EasyOCR failed on page {page_num} of {file_path}: {e}")
            except TimeoutError as e:
                logger.warning(f"{e}. Returning partial OCR content for {file_name}.")
            except Exception as e:
                logger.info(f"OCR extraction error: {e}")

        if not all_docs:
            logger.warning(f"No content extracted from {file_path}")
        logger.info(
            f"Finished PDF: {file_name} -> {len(all_docs)} chunks in {time.time() - started:.1f}s"
        )
        return all_docs


class DeferredExtractor:
    def __init__(self, factory):
        self._factory = factory
        self._extractor = None

    def load_data(self, *args, **kwargs):
        if self._extractor is None:
            self._extractor = self._factory()
        return self._extractor.load_data(*args, **kwargs)


class DoclingReader:
    """Parse modern Office / structured documents via Docling and emit markdown
    chunks suitable for vector embeddings. Supports the Docling-native formats
    (.docx, .pptx, .xlsx, .html, .md, .csv, etc.); legacy binary formats
    (.doc/.xls/.ppt) and .rtf must be routed to a different reader. Falls back
    to UnstructuredReader when docling can't recognize the file (e.g. a legacy
    .doc renamed with a .docx extension)."""

    _converter = None
    _fallback = None

    @classmethod
    def _get_converter(cls):
        if cls._converter is None:
            from docling.document_converter import DocumentConverter
            cls._converter = DocumentConverter()
        return cls._converter

    @classmethod
    def _get_fallback(cls):
        if cls._fallback is None:
            cls._fallback = UnstructuredReader()
        return cls._fallback

    def _fallback_load(self, file_path_str, file_name, reason):
        logger.info(
            f"Docling could not parse {file_name} ({reason}); falling back to UnstructuredReader."
        )
        try:
            fb_docs = self._get_fallback().load_data(file_path_str) or []
        except Exception as e:
            logger.error(f"UnstructuredReader fallback failed for {file_path_str}: {e}")
            return []
        out: List[Document] = []
        for fb in fb_docs:
            text = normalize_text(getattr(fb, "text", "") or "")
            if not text:
                continue
            for chunk_index, chunk in enumerate(
                chunk_text(text, chunk_size=800, chunk_overlap=160)
            ):
                out.append(Document(text=chunk, metadata={
                    "file_path": file_path_str,
                    "file_name": file_name,
                    "chunk_index": chunk_index,
                    "extracted_by": "unstructured_fallback",
                }))
        return out

    def load_data(self, file_path, **kwargs):
        file_path_str = str(file_path)
        file_name = os.path.basename(file_path_str)
        try:
            result = self._get_converter().convert(file_path_str)
            text = normalize_text(result.document.export_to_markdown())
        except Exception as e:
            return self._fallback_load(file_path_str, file_name, str(e))
        if not text:
            return self._fallback_load(file_path_str, file_name, "empty docling output")
        docs: List[Document] = []
        # ~800-word chunks with 160-word overlap fits comfortably in bge-m3's
        # 8192-token context window (~6x headroom) while keeping per-section
        # semantic focus. Going larger would dilute retrieval ranking; going
        # smaller would split section headers from their items. Send bigger
        # context to the LLM downstream via JARVIS_SOURCE_HINT_CHUNK_CHARS,
        # JARVIS_MAX_EVIDENCE_LINES, and JARVIS_RETRIEVAL_TOP_K instead.
        for chunk_index, chunk in enumerate(
            chunk_text(text, chunk_size=800, chunk_overlap=160)
        ):
            docs.append(Document(text=chunk, metadata={
                "file_path": file_path_str,
                "file_name": file_name,
                "chunk_index": chunk_index,
                "extracted_by": "docling",
            }))
        return docs


class MHTMLReader:
    """Parse .mhtml/.mht MIME web-archive files (single-file "webpage, complete"
    saves). The container is a MIME multipart message; extract its primary
    text/html part with the stdlib ``email`` parser, then delegate the HTML to
    ``html_reader`` (Docling, with its UnstructuredReader fallback) so extraction
    stays identical to a plain .html file."""

    def __init__(self, html_reader):
        self._html_reader = html_reader

    @staticmethod
    def _extract_html(file_path_str: str) -> Optional[str]:
        import email
        from email import policy

        with open(file_path_str, "rb") as handle:
            message = email.message_from_binary_file(handle, policy=policy.default)
        html_part = None
        text_part = None
        for part in message.walk():
            content_type = part.get_content_type()
            if content_type == "text/html" and html_part is None:
                html_part = part
            elif content_type == "text/plain" and text_part is None:
                text_part = part
        chosen = html_part or text_part
        if chosen is None:
            return None
        content = chosen.get_content()
        return content if isinstance(content, str) else str(content)

    def load_data(self, file_path, **kwargs):
        file_path_str = str(file_path)
        file_name = os.path.basename(file_path_str)
        try:
            html = self._extract_html(file_path_str)
        except Exception as e:
            logger.error(f"MHTMLReader could not parse {file_name}: {e}")
            return []
        if not html:
            logger.info(f"MHTMLReader found no text/html or text/plain part in {file_name}.")
            return []
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(
                "w", suffix=".html", delete=False, encoding="utf-8"
            ) as tmp:
                tmp.write(html)
                tmp_path = tmp.name
            docs = self._html_reader.load_data(tmp_path, **kwargs) or []
        except Exception as e:
            logger.error(f"MHTMLReader failed to extract HTML from {file_name}: {e}")
            return []
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.remove(tmp_path)
        # Rewrite provenance from the temp .html back to the original archive.
        for doc in docs:
            doc.metadata["file_path"] = file_path_str
            doc.metadata["file_name"] = file_name
        return docs


class EmlxReader:
    """Parse an Apple Mail ``.emlx`` message into a single text Document.

    ``.emlx`` framing is an ASCII decimal byte-count line, the raw RFC-822
    message of exactly that many bytes, then an Apple metadata plist. Parse the
    message span with the stdlib ``email`` parser (same approach as
    ``MHTMLReader``), preferring the ``text/plain`` part and falling back to a
    tag-stripped ``text/html`` part. From/To/Cc/Date/Subject headers are kept as
    metadata so an answer can attribute a quote to a specific message."""

    _TAG_RE = re.compile(r"<[^>]+>")
    # <style>/<script> (and CSS/JS text between the tags) must be removed
    # wholesale: stripping only the tags would leave the rule/code text as body
    # content. DOTALL so a block spanning newlines is matched; the backreference
    # pairs each opening tag with its own closing tag.
    _STYLE_SCRIPT_RE = re.compile(
        r"<(style|script)\b[^>]*>.*?</\1\s*>", re.IGNORECASE | re.DOTALL
    )
    # HTML comments can also wrap content (e.g. conditional comments in mail).
    _COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
    # Zero-width / BOM characters that survive entity decoding and add no meaning.
    _ZERO_WIDTH_RE = re.compile(r"[\ufeff\u200b\u200c\u200d\u2060]")
    _WHITESPACE_RE = re.compile(r"\s+")

    @classmethod
    def _html_to_text(cls, html: str) -> str:
        """Reduce an HTML email body to readable text: drop <style>/<script>
        blocks (and their contents) and comments, strip remaining tags, decode
        entities, remove zero-width characters, and collapse whitespace."""
        without_blocks = cls._STYLE_SCRIPT_RE.sub(" ", html)
        without_comments = cls._COMMENT_RE.sub(" ", without_blocks)
        without_tags = cls._TAG_RE.sub(" ", without_comments)
        decoded = html_lib.unescape(without_tags)
        decoded = cls._ZERO_WIDTH_RE.sub("", decoded)
        return cls._WHITESPACE_RE.sub(" ", decoded).strip()

    @staticmethod
    def _message_bytes(raw: bytes) -> bytes:
        newline = raw.find(b"\n")
        if newline == -1:
            return raw
        try:
            count = int(raw[:newline].strip())
        except ValueError:
            return raw  # not framed as expected; parse the whole payload
        start = newline + 1
        return raw[start:start + count]

    @classmethod
    def _body_text(cls, message) -> str:
        plain = html = None
        for part in message.walk():
            ctype = part.get_content_type()
            if ctype == "text/plain" and plain is None:
                try:
                    plain = part.get_content()
                except Exception:
                    plain = None
            elif ctype == "text/html" and html is None:
                try:
                    html = part.get_content()
                except Exception:
                    html = None
        if isinstance(plain, str) and plain.strip():
            return plain
        if isinstance(html, str) and html.strip():
            return cls._html_to_text(html)
        return ""

    def load_data(self, file_path, **kwargs):
        import email
        from email import policy

        file_path_str = str(file_path)
        file_name = os.path.basename(file_path_str)
        try:
            with open(file_path_str, "rb") as handle:
                raw = handle.read()
            message = email.message_from_bytes(self._message_bytes(raw), policy=policy.default)
        except Exception as e:
            logger.error(f"EmlxReader could not parse {file_name}: {e}")
            return []

        headers = {
            "email_from": str(message.get("From", "") or ""),
            "email_to": str(message.get("To", "") or ""),
            "email_cc": str(message.get("Cc", "") or ""),
            "email_date": str(message.get("Date", "") or ""),
            "email_subject": str(message.get("Subject", "") or ""),
        }
        body = self._body_text(message)
        header_block = "\n".join(
            f"{label}: {headers[key]}"
            for label, key in (
                ("From", "email_from"), ("To", "email_to"), ("Cc", "email_cc"),
                ("Date", "email_date"), ("Subject", "email_subject"),
            )
            if headers[key]
        )
        text = f"{header_block}\n\n{body}".strip()
        if not text:
            logger.info(f"EmlxReader found no readable content in {file_name}.")
            return []
        metadata = {"file_path": file_path_str, "file_name": file_name, "extracted_by": "emlx"}
        metadata.update({k: v for k, v in headers.items() if v})
        return [Document(text=text, metadata=metadata)]


def staged_email_attachments(staged_path: str) -> list[tuple[str, str]]:
    """Resolve a staged email's attachments to local files, on demand.

    The staged ``.emlx`` is the pointer: its MIME headers list each attachment
    (``Content-Disposition: attachment``; inline logos and tracking pixels are
    skipped) even when Apple Mail stored the bytes outside the message. External
    files come from the attachment mirror that mail refresh keeps beside the
    corpus; an attachment embedded in the message itself is written into that
    mirror the first time it is needed. Nothing is extracted or indexed here.
    Returns ``(attachment name, local path)`` pairs; unreadable input yields []."""
    import email
    from email import policy

    from jarvisrag.docsync import mail_attachment_dir

    stem = os.path.basename(staged_path).split(".", 1)[0]
    mirror = os.path.join(mail_attachment_dir(os.path.dirname(os.path.abspath(staged_path))), stem)
    try:
        with open(staged_path, "rb") as handle:
            message = email.message_from_bytes(
                EmlxReader._message_bytes(handle.read()), policy=policy.default
            )
    except Exception:  # noqa: BLE001 - a bad pointer just yields no attachments
        return []

    mirrored: dict[str, str] = {}
    if os.path.isdir(mirror):
        for root, _dirs, files in os.walk(mirror):
            for fname in files:
                if not fname.startswith(".") and not fname.endswith(".mirror-tmp"):
                    mirrored.setdefault(fname.lower(), os.path.join(root, fname))

    found: list[tuple[str, str]] = []
    used: set[str] = set()
    for part in message.walk():
        if part.is_multipart() or part.get_content_disposition() != "attachment":
            continue
        name = os.path.basename(str(part.get_filename() or "").replace("\\", "/")).strip()
        if not name:
            continue
        try:
            payload = part.get_payload(decode=True) or b""
        except Exception:  # noqa: BLE001
            payload = b""
        path = None
        if payload:
            path = os.path.join(mirror, "embedded", name)
            try:
                if not (os.path.exists(path) and os.path.getsize(path) == len(payload)):
                    os.makedirs(os.path.dirname(path), exist_ok=True)
                    with open(path, "wb") as out:
                        out.write(payload)
            except OSError:
                path = None
        else:
            path = mirrored.get(name.lower())
        if path and path not in used:
            used.add(path)
            found.append((name, path))
    # Apple Mail may store a file under a slightly different name than the MIME
    # header; keep any mirrored file that no header matched rather than lose it.
    for path in mirrored.values():
        if path not in used:
            used.add(path)
            found.append((os.path.basename(path), path))
    return found


jpeg_extractor = DeferredExtractor(JpegExtractor)
docling_reader = DoclingReader()
mhtml_reader = MHTMLReader(docling_reader)
file_extractor = {
    ".pdf": CombinedPDFReader(),
    ".docx": docling_reader,
    ".doc": UnstructuredReader(),
    ".rtf": UnstructuredReader(),
    ".csv": UnstructuredReader(),
    ".txt": None,
    ".html": docling_reader,
    ".htm": docling_reader,
    ".mhtml": mhtml_reader,
    ".mht": mhtml_reader,
    ".emlx": EmlxReader(),
    ".jpeg": jpeg_extractor,
    ".jpg": jpeg_extractor,
    ".png": jpeg_extractor,
    ".gif": jpeg_extractor,
    ".tif": jpeg_extractor,
    ".tiff": jpeg_extractor,
    ".bmp": jpeg_extractor,
    ".xlsx": docling_reader,
    ".xls": ExcelReader(),
    ".pptx": docling_reader,
    ".ppt": UnstructuredReader(),
}
file_extractor[".zip"] = ZipExtractor(file_extractor)


def has_files_with_extensions(directory, extensions):
    if not os.path.exists(directory):
        return False
    for root, _, files in os.walk(directory):
        for file in files:
            if any(file.lower().endswith(ext) for ext in extensions):
                return True
    return False


# Bump whenever extraction OUTPUT changes in a way not captured by
# _extraction_mode_signature() (e.g. VLM dtype/preprocessing/generation changes
# in jarvisrag/extraction/vlm.py). v3: SmolVLM2 float32 + EXIF transpose +
# downscale + repetition controls. v4: Apple Vision OCR backend on the text path.
# v5: Apple Vision OCR row-band assembly (preserves table rows/columns).
# v6: document-level OCR validator + escalation on whole-page failure.
# v7: HTML (.html/.htm via Docling) + MHTML (.mhtml/.mht) extraction support.
EXTRACTION_CACHE_VERSION = 7


# Bump when EMAIL (.emlx) extraction OUTPUT changes without a global
# EXTRACTION_CACHE_VERSION bump, so ONLY email caches invalidate (a global bump
# would needlessly re-OCR every PDF/image). v1: HTML body cleaning -- strip
# <style>/<script> and their contents, decode entities, drop zero-width chars,
# collapse whitespace.
EMAIL_EXTRACTION_VERSION = 1
_EMAIL_EXTENSIONS = (".emlx",)


def _extractor_signature(file_path: str) -> str:
    """A per-file extractor signature that varies only for file types whose
    reader output can change independently of the global cache version. Empty
    for every other file so unrelated caches are never invalidated."""
    if file_path.lower().endswith(_EMAIL_EXTENSIONS):
        return f"emlx{EMAIL_EXTRACTION_VERSION}"
    return ""


def _extraction_cache_path(cache_dir: str, file_path: str) -> str:
    key = hashlib.sha256(os.path.abspath(file_path).encode("utf-8")).hexdigest()
    return os.path.join(cache_dir, f"{key}.json")


def _extraction_mode_signature() -> str:
    """Signature of the settings that change extraction OUTPUT.

    Folded into the cache validation so that toggling the page pipeline, Surya,
    or the OCR/native-text thresholds automatically invalidates stale cached
    extractions (otherwise a prior run's docs would be replayed under a new
    mode, masking the change).
    """
    return (
        f"pipeline={int(JARVIS_ENABLE_PAGE_PIPELINE)};"
        f"surya={int(JARVIS_ENABLE_SURYA)};"
        f"minconf={JARVIS_OCR_MIN_CONF};"
        f"minchars={JARVIS_PDF_MIN_CHARS_PER_PAGE};"
        f"dpi={JARVIS_PDF_OCR_DPI};"
        f"vlm={int(JARVIS_ENABLE_VLM)};"
        f"vlmmodel={JARVIS_VLM_MODEL if JARVIS_ENABLE_VLM else ''};"
        f"vlmcaption={int(JARVIS_VLM_CAPTION)};"
        f"appleocr={int(JARVIS_ENABLE_APPLE_OCR)};"
        f"applelangs={JARVIS_APPLE_OCR_LANGS if JARVIS_ENABLE_APPLE_OCR else ''};"
        f"applelevel={JARVIS_APPLE_OCR_LEVEL if JARVIS_ENABLE_APPLE_OCR else ''};"
        f"ocrescalate={int(JARVIS_OCR_ESCALATE_ON_FAILURE)}"
    )


def _serialize_documents_for_cache(docs: List[Document]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for doc in docs:
        metadata = doc.metadata or {}
        safe_meta: Dict[str, Any] = {}
        for k, v in metadata.items():
            if isinstance(v, os.PathLike):
                safe_meta[str(k)] = os.fspath(v)
            elif isinstance(v, (str, int, float, bool)) or v is None:
                safe_meta[str(k)] = v
            else:
                safe_meta[str(k)] = str(v)
        out.append({
            "id_": doc.id_,
            "text": doc.text or "",
            "metadata": safe_meta,
        })
    return out


def _deserialize_documents_from_cache(entries: List[Dict[str, Any]]) -> List[Document]:
    docs: List[Document] = []
    for entry in entries or []:
        doc = Document(
            text=entry.get("text", "") or "",
            metadata=dict(entry.get("metadata") or {}),
        )
        cached_id = entry.get("id_")
        if cached_id:
            doc.id_ = cached_id
        docs.append(doc)
    return docs


def _walk_files(input_dir: str, extensions: List[str], recursive: bool = True) -> List[str]:
    matched: List[str] = []
    if not os.path.exists(input_dir):
        return matched
    lowered = tuple(ext.lower() for ext in extensions)
    if recursive:
        for root, _, files in os.walk(input_dir):
            for file in files:
                if file.lower().endswith(lowered):
                    matched.append(os.path.join(root, file))
    else:
        for name in os.listdir(input_dir):
            full = os.path.join(input_dir, name)
            if os.path.isfile(full) and name.lower().endswith(lowered):
                matched.append(full)
    matched.sort()
    return matched


def load_documents_with_cache(
    input_dir: str,
    required_exts: List[str],
    file_extractor: Dict[str, Any],
    cache_dir: str,
    recursive: bool = True,
) -> List[Document]:
    """Extract documents from a directory with per-file (path, mtime, size) caching.

    On cache hit, the original Documents are restored from disk without
    re-running OCR / PDF parsing. On miss or stale (mtime or size changed),
    the file is re-extracted via SimpleDirectoryReader using the provided
    file_extractor mapping, and the cache entry is rewritten.

    Set JARVIS_DISABLE_EXTRACTION_CACHE=1 to bypass the cache.
    """
    from llama_index.core import SimpleDirectoryReader  # local import to avoid cycles

    if os.environ.get("JARVIS_DISABLE_EXTRACTION_CACHE", "0") == "1":
        reader = SimpleDirectoryReader(
            input_dir=input_dir,
            recursive=recursive,
            required_exts=required_exts,
            file_extractor=file_extractor,
        )
        return reader.load_data()

    os.makedirs(cache_dir, exist_ok=True)
    files = _walk_files(input_dir, required_exts, recursive=recursive)
    if not files:
        return []

    all_docs: List[Document] = []
    cache_hits = 0
    cache_misses = 0
    extract_failures = 0
    mode_sig = _extraction_mode_signature()

    for file_path in files:
        try:
            stat = os.stat(file_path)
        except OSError as exc:
            print(f"[jarvis] Skipping unreadable file {file_path}: {exc}")
            continue
        mtime_ns = getattr(stat, "st_mtime_ns", int(stat.st_mtime * 1_000_000_000))
        size = stat.st_size
        extractor_sig = _extractor_signature(file_path)
        cache_path = _extraction_cache_path(cache_dir, file_path)
        cached: Optional[Dict[str, Any]] = None
        if os.path.isfile(cache_path):
            try:
                with open(cache_path, "r", encoding="utf-8") as fh:
                    cached = json.load(fh)
            except Exception:
                cached = None

        if (
            cached
            and cached.get("version") == EXTRACTION_CACHE_VERSION
            and cached.get("extraction_mode") == mode_sig
            and (not extractor_sig or cached.get("extractor_sig") == extractor_sig)
            and cached.get("mtime_ns") == mtime_ns
            and cached.get("size") == size
            and cached.get("file_path") == os.path.abspath(file_path)
        ):
            docs = _deserialize_documents_from_cache(cached.get("documents") or [])
            all_docs.extend(docs)
            cache_hits += 1
            continue

        cache_misses += 1
        if JARVIS_DEBUG:
            if cached is None:
                miss_reason = "no cache file"
            else:
                reasons = []
                if cached.get("version") != EXTRACTION_CACHE_VERSION:
                    reasons.append(f"version({cached.get('version')}!={EXTRACTION_CACHE_VERSION})")
                if cached.get("extraction_mode") != mode_sig:
                    reasons.append(f"mode({cached.get('extraction_mode')}!={mode_sig})")
                if extractor_sig and cached.get("extractor_sig") != extractor_sig:
                    reasons.append(f"extractor({cached.get('extractor_sig')}!={extractor_sig})")
                if cached.get("mtime_ns") != mtime_ns:
                    reasons.append(f"mtime({cached.get('mtime_ns')}!={mtime_ns})")
                if cached.get("size") != size:
                    reasons.append(f"size({cached.get('size')}!={size})")
                if cached.get("file_path") != os.path.abspath(file_path):
                    reasons.append(f"path({cached.get('file_path')}!={os.path.abspath(file_path)})")
                miss_reason = ",".join(reasons) or "unknown"
            print(f"[debug] extraction-cache MISS {os.path.basename(file_path)}: {miss_reason}")
        try:
            single_file_extractor = {
                ext: reader for ext, reader in (file_extractor or {}).items() if reader is not None
            }
            reader = SimpleDirectoryReader(
                input_files=[file_path],
                file_extractor=single_file_extractor,
            )
            docs = reader.load_data()
        except Exception as exc:
            extract_failures += 1
            print(f"[jarvis] Extraction failed for {file_path}: {exc}")
            continue

        all_docs.extend(docs)
        payload = {
            "version": EXTRACTION_CACHE_VERSION,
            "extraction_mode": mode_sig,
            "extractor_sig": extractor_sig,
            "file_path": os.path.abspath(file_path),
            "mtime_ns": mtime_ns,
            "size": size,
            "documents": _serialize_documents_for_cache(docs),
        }
        try:
            with open(cache_path, "w", encoding="utf-8") as fh:
                json.dump(payload, fh)
        except Exception as exc:
            print(f"[jarvis] Could not write extraction cache for {file_path}: {exc}")

    print(
        f"[jarvis] Extraction cache: {cache_hits} hit, {cache_misses} miss"
        + (f", {extract_failures} failed" if extract_failures else "")
        + f" (out of {len(files)} files)."
    )
    return all_docs
